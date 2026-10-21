"""Build the T5 evidence report from recorded physical and software evidence."""
from pathlib import Path
from datetime import datetime
import json
import shutil
from zoneinfo import ZoneInfo

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
EV = HERE / 'evidencias'
LOCAL = ROOT / '.local/taller5-documento'
DIAG = HERE / 'diagramas'
DIAG.mkdir(exist_ok=True)
LOCAL.mkdir(parents=True, exist_ok=True)
TZ = ZoneInfo('America/Tijuana')
physical = json.loads((ROOT / 'docs/evidence/public-flow-esp32-android-2026-10-21.json').read_text())
tds = json.loads((ROOT / 'docs/evidence/tds-events-2026-10-21.json').read_text())
api = json.loads((EV / 'api-publica.json').read_text())
records = json.loads((EV / 'registros-publicos.json').read_text())['rows']
photos = json.loads((HERE/'fotos.json').read_text(encoding='utf8')) if (HERE/'fotos.json').exists() else []


def font(size=30, bold=False):
    return ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf' if bold else 'C:/Windows/Fonts/arial.ttf', size)


def make_diagrams():
    # These are diagrams of the implemented flow, never photographs of the bench.
    for kind in ('flujo', 'conexiones'):
        image = Image.new('RGB', (1680, 650), 'white')
        draw = ImageDraw.Draw(image)
        svg = ['<svg xmlns="http://www.w3.org/2000/svg" width="1680" height="650" viewBox="0 0 1680 650"><rect width="1680" height="650" fill="white"/>']
        def text(x,y,value,size=28,bold=False):
            draw.text((x,y),value,font=font(size,bold),fill='#162b35')
            from html import escape
            svg.append(f'<text x="{x}" y="{y+size}" font-family="Arial" font-size="{size}" font-weight="{700 if bold else 400}" fill="#162b35">{escape(value)}</text>')
        def box(x,y,w,h,label,sub=''):
            draw.rounded_rectangle((x,y,x+w,y+h),radius=12,fill='#edf4f7',outline='#9eafb8',width=2)
            svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="12" fill="#edf4f7" stroke="#9eafb8" stroke-width="2"/>')
            text(x+18,y+20,label,29,True)
            if sub:text(x+18,y+66,sub,23)
        def arrow(x,y,x2,y2):
            import math
            draw.line((x,y,x2,y2),fill='#36596b',width=4)
            a=math.atan2(y2-y,x2-x)
            pts=[(x2,y2),(x2-16*math.cos(a-.45),y2-16*math.sin(a-.45)),(x2-16*math.cos(a+.45),y2-16*math.sin(a+.45))]
            draw.polygon(pts,fill='#36596b')
            svg.append(f'<line x1="{x}" y1="{y}" x2="{x2}" y2="{y2}" stroke="#36596b" stroke-width="4"/>')
            svg.append('<polygon points="'+' '.join(f'{a:.1f},{b:.1f}' for a,b in pts)+'" fill="#36596b"/>')
        if kind=='flujo':
            text(30,12,'Recorrido implementado de una lectura real',36,True)
            box(30,100,290,120,'Sensores físicos','Luz, nivel, pH, TDS, °C')
            box(400,100,250,120,'ESP32','Wi-Fi y telemetría v3')
            box(745,100,375,120,'Mosquitto en la VPS','MQTT TLS público 8884')
            box(1220,100,400,120,'Bridge Python','Recibe y envía a Laravel')
            arrow(320,160,400,160);arrow(650,160,745,160);arrow(1120,160,1220,160)
            box(1220,340,400,120,'Laravel en la VPS','Valida y confirma guardado')
            arrow(1420,220,1420,340)
            box(745,340,375,120,'PostgreSQL','Red interna y volumen propio')
            arrow(1220,390,1120,390);arrow(1120,425,1220,425)
            box(30,340,620,120,'Android en el teléfono','Consulta Laravel por HTTPS 9443')
            # GET travels to Laravel; the response contains persisted measurements.
            draw.line((1420,460,1420,540,350,540,350,460),fill='#36596b',width=4)
            svg.append('<path d="M1420 460V540H350V460" stroke="#36596b" stroke-width="4" fill="none"/>')
            arrow(350,505,350,460)
            text(455,565,'Respuesta de la API con las mediciones almacenadas',27)
        else:
            text(30,12,'Conexiones de señal utilizadas por el firmware',36,True)
            box(30,90,400,110,'DS18B20','Datos a GPIO 4')
            box(30,250,400,110,'Flotador','Contacto a GPIO 27 y GND')
            box(30,410,400,110,'BH1750','I2C compartido')
            box(570,220,460,155,'ESP32','GPIO 21 SDA y GPIO 22 SCL')
            arrow(430,145,570,255);arrow(430,305,570,300);arrow(430,465,570,345)
            box(1170,220,470,155,'ADS1115 dirección 0x48','A0 pH     A1 TDS')
            arrow(1030,295,1170,295)
            box(1170,420,470,110,'Módulos analógicos','pH a A0     TDS a A1')
            arrow(1400,420,1400,375)
            text(35,575,'Diagrama de señales según código. Verificar alimentación y masa común en el montaje.',27)
        svg.append('</svg>')
        (DIAG/f'{kind}.svg').write_text('\n'.join(svg),encoding='utf8')
        image.save(DIAG/f'{kind}.png')


make_diagrams()
doc = Document()
sec = doc.sections[0]
sec.page_width,sec.page_height=Inches(8.5),Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.7)
for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
    style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.space_after=Pt(7)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(27)
doc.styles['Heading 1'].font.size=Pt(20)
doc.styles['Heading 2'].font.size=Pt(13)
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].font.italic=False
# The bundled Word template carries a blue Title paragraph border. Remove it.
for border in list(doc.styles.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('HydroMate  |  Taller 5  |  ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title='HydroMate Evidencias del Taller 5'
doc.core_properties.subject='Prueba de viabilidad con sensores físicos y backend en la nube'
doc.core_properties.author='HydroMate'
doc.core_properties.comments=''


def p(text='',bold=False,style=None):
    para=doc.add_paragraph(style=style)
    run=para.add_run(text);run.bold=bold
    return para


def page(title):
    doc.add_page_break();doc.add_heading(title,level=1)


def table(headers,rows,widths=None,size=10):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths:
        for col,w in zip(t.columns,widths):col.width=Inches(w)
    for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
    for values in rows:
        for c,v in zip(t.add_row().cells,values):c.text=str(v)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for idx,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        for j,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[j])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr();margin=OxmlElement('w:tcMar')
            for edge in ('top','bottom','left','right'):
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'70');el.set(qn('w:type'),'dxa');margin.append(el)
            tcpr.append(margin)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCEAF2' if idx==0 else ('F5F8FA' if idx%2==0 else 'FFFFFF'));tcpr.append(shade)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2);para.paragraph_format.line_spacing=1.03
                if idx==0:para.paragraph_format.keep_with_next=True
                for run in para.runs:run.font.size=Pt(size);run.bold=idx==0;run.font.color.rgb=RGBColor(0,0,0)
    p().paragraph_format.space_after=Pt(0)
    return t


def code(text):
    para=p(text)
    para.paragraph_format.line_spacing=1.0;para.paragraph_format.space_after=Pt(9)
    for r in para.runs:r.font.name='Consolas';r.font.size=Pt(9)


def figure(path,caption,width=6.9):
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(path),width=Inches(width))
    para.paragraph_format.keep_with_next=True
    p(caption,style='Caption')


def localtime(value):
    return datetime.fromisoformat(value.replace('Z','+00:00')).astimezone(TZ).strftime('%H:%M:%S')


doc.add_paragraph('HydroMate',style='Title')
doc.add_heading('Evidencias del Taller 5',level=1)
p('Prueba de viabilidad de la arquitectura',style='Subtitle')
p('Corte de evidencias 21 de octubre de 2026\nEntrega final del proyecto 2 de diciembre de 2026')
p('La prueba ya recorre el camino que pide el taller: una lectura física sale del ESP32, llega al backend de la VPS, queda guardada en PostgreSQL y se puede consultar desde la aplicación Android. La idea de esta entrega es mostrar ese recorrido y señalar exactamente con qué se demuestra cada punto.')
figure(DIAG/'flujo.png','Figura 1. Flujo implementado. El teléfono consulta a Laravel; nunca entra directamente a PostgreSQL.',6.95)
p('Qué se está demostrando',bold=True)
p('Se utilizan temperatura, iluminación, nivel de agua, pH y TDS reales. La prueba es de comunicación y almacenamiento; no demuestra todavía calibración química ni control automático de dosificación.')
table(['Elemento','Resultado que se puede revisar'],[
('Recorrido físico público','14 transmisiones nuevas consecutivas confirmadas y comparadas con la API.'),
('Lecturas actuales','Temperatura y TDS reales visibles en Android; historial consultado por HTTPS.'),
('Planeación','Trabajo restante hasta el 2 de diciembre, presupuesto preliminar y matriz de riesgos.'),
('Control de actuadores','Aplazamiento hasta alfa, autorizado según confirmación del responsable.')],[1.5,5.5],10.5)

page('Montaje y sensores utilizados')
p('El montaje usa un ESP32 como sistema embebido. Los sensores digitales entregan su lectura por GPIO o I2C; pH y TDS pasan por el ADS1115, que convierte sus señales analógicas a valores que el programa puede procesar.')
figure(DIAG/'conexiones.png','Figura 2. Conexiones de señal utilizadas. Este esquema no sustituye la revisión eléctrica del montaje.',6.95)
table(['Entrada','Conexión utilizada','Dato enviado'],[
('DS18B20','GPIO 4; pull-up de 4.7 kΩ documentado','Temperatura en °C'),
('BH1750','I2C SDA 21 y SCL 22','Iluminación en lux'),
('Flotador','GPIO 27 con INPUT_PULLUP','true agua; false nivel bajo'),
('pH','Salida analógica a ADS1115 A0','pH sin unidad'),
('TDS de reemplazo','Salida analógica a ADS1115 A1','TDS en ppm'),
('ADS1115','I2C compartido, dirección 0x48','Conversión de pH y TDS')],[1.45,3.55,2],10.5)
p('Se conserva masa común y se revisa la alimentación antes de montar la versión definitiva. La correspondencia A1 del TDS y GPIO 4 de temperatura fue confirmada durante la integración. EC permanece sin sensor independiente y se envía como no disponible.')
if photos:
    p('Las figuras 5 y 6 muestran las sondas en agua y el armado real en protoboard. Complementan este diagrama de señales con fotografías del banco utilizado.')
p('Código relacionado: embedded/HydroMateTelemetry/HydroMateTelemetry.ino, sensor_config.h y telemetry_sensors.h.',style='Caption')

page('Cómo lee y transmite el ESP32')
p('El firmware lee las entradas reales y prepara un mensaje JSON de telemetría v3. La temperatura se obtiene sin detener la conversión completa; en las entradas analógicas se usa una mediana de nueve lecturas para reducir valores aislados. Después se aplican las conversiones de pH y TDS y se redondea a dos decimales.')
p('TDS se compensa con la temperatura medida por DS18B20. Cuando un sensor es inválido, su campo queda en null con un origen explícito; no se sustituye por una lectura inventada. Las lecturas actuales de TDS usan la temperatura real disponible.')
table(['Función mínima','Cómo se cumple'],[
('Entrada física','DS18B20, BH1750, flotador, pH y TDS del montaje.'),
('Procesamiento básico','Mediana, conversión a unidades, validación de rango y redondeo.'),
('Conexión a Internet','Wi-Fi de 2.4 GHz y sincronización de reloj para TLS.'),
('Envío al backend','MQTT TLS a hydromate-kevin.duckdns.org:8884; el bridge entrega a Laravel.'),
('Identificación','device_id hydromate-01, sequence y captured_at en UTC. Laravel agrega created_at.')],[1.65,5.35],10.5)
p('Además del respaldo de 30 segundos, se envía una muestra cuando cambia claramente una variable. El flotador tiene antirrebote de 150 ms. Los umbrales están concentrados en sensor_config.h; por ejemplo, luz cambia por 50 lux o 30 %, y temperatura por 1 °C.')
p('Cada muestra pendiente se guarda en NVS, una memoria que conserva información aunque se reinicie el ESP32. El dispositivo reenvía el mismo mensaje hasta recibir STORED, la confirmación posterior al guardado en la base de datos. Así se evita confundir un reintento con una medición nueva.')
code('hydromate/hydromate-01/telemetry  →  lectura JSON\nhydromate/hydromate-01/ack        ←  confirmación stored')
p('Wi-Fi, usuario MQTT, contraseña y CA se configuran en secrets.h privado, excluido de Git. Producción valida el certificado y el hostname; no utiliza setInsecure.')

page('Backend desplegado en la nube')
p('El backend es una API Laravel desplegada en una VPS Ubuntu de OVHcloud. Mosquitto recibe la telemetría y el bridge Python la convierte en una petición HTTP interna. Laravel valida el contenido y lo guarda en PostgreSQL.')
p('La API se consulta públicamente en https://hydromate-kevin.duckdns.org:9443. En la revisión del 21 de octubre, /up, el historial y el estado de hydromate-01 respondieron HTTP 200. Se recuperaron 20 registros con origen real y el estado de telemetría fue recent.')
table(['Operación','Resultado y alcance'],[
('POST /api/measurements','Uso interno del bridge. 201 si registra, 422 si el dato es incorrecto y 409 si la identidad ya existe.'),
('GET /api/measurements','Consulta historial por device_id y limit; real_only=1 filtra datos simulados.'),
('GET /api/measurements/latest','Recupera la última medición almacenada.'),
('GET /api/devices/hydromate-01/status','Informa antigüedad de telemetría: no_data, recent o stale.'),
('GET /up','Comprobación HTTP del servicio; la evidencia pública respondió 200.')],[2.65,4.35],10.5)
p('Los campos indispensables incluyen identidad y secuencia, versión, valores y origen de cada sensor. El flotador debe ser un booleano JSON y los números deben estar dentro del rango aceptado. El backend conserva el mismo contrato para registrar y consultar.')
p('El bridge no confirma STORED hasta que Laravel registra la medición o verifica que ya existía idéntica. Su bandeja SQLite es una ayuda de entrega; PostgreSQL sigue siendo la base principal de mediciones.')
p('Despliegue y separación',bold=True)
p('Docker rootless ejecuta únicamente los servicios de HydroMate. API interna en 127.0.0.1:8081; HTTPS público en 9443; MQTT TLS en 8884. PostgreSQL no publica 5432 al host. Los volúmenes, red y secretos pertenecen a HydroMate. Android se ejecuta en el teléfono.')
p('Evidencia: evidencias/api-publica.json; compose.production.yml; deployment/deploy.sh; docs/PUBLIC_ACCESS.md. La comprobación pública fue de solo lectura y no requirió acceso SSH.',style='Caption')
p('Código: hydromate-backend/app/Http/Requests/StoreMeasurementRequest.php valida; app/Http/Controllers/Api/MeasurementController.php registra y consulta; scripts/mqtt_bridge.py recibe MQTT y entrega a la API. Las rutas app/ parten de hydromate-backend/.',style='Caption')

page('Estructura de la base de datos')
p('La colección mínima se resuelve con la tabla measurements. Cada fila representa una adquisición de una torre y contiene las variables disponibles en ese instante. Las unidades se conservan en el contrato y en los nombres de campo; no hace falta una tabla por sensor.')
table(['Campo','Qué identifica o almacena'],[
('id','Identificador de la fila en PostgreSQL.'),
('device_id y sequence','Torre y número de adquisición; pareja única para evitar duplicados.'),
('temperature_c y light_lux','Temperatura en °C e iluminación en lux.'),
('water_present','Estado discreto del flotador, sin unidad.'),
('ph y tds_ppm','Valor de pH y concentración aproximada en ppm.'),
('ec_us_cm','Conductividad independiente en µS/cm; actualmente null/unavailable.'),
('sources','Origen por variable: real, simulated, unavailable o invalid.'),
('captured_at y created_at','Fecha/hora de adquisición y de recepción, en UTC.'),
('message_version y reason','Versión del contrato y motivo de la adquisición.')],[2.35,4.65],10.5)
p('Consulta de recuperación',bold=True)
code("SELECT id, device_id, sequence, temperature_c,\n       ph, tds_ppm, light_lux, water_present,\n       captured_at, created_at\nFROM measurements\nWHERE device_id = 'hydromate-01'\nORDER BY created_at DESC, id DESC\nLIMIT 20;")
p('Android usa el GET de Laravel para obtener esta información; no ejecuta SQL ni recibe credenciales de PostgreSQL. El esquema se conserva en las migraciones, incluida la ampliación v3 de octubre que permite valores null y hora de adquisición.')
p('Fuentes: hydromate-backend/database/migrations/ y pruebas/esquema-postgresql-2026-10-07.json. El esquema exportado es un antecedente local; las migraciones v3 describen las columnas añadidas posteriormente.',style='Caption')

page('Diez registros reales recuperados')
p('Esta tabla contiene diez filas obtenidas del backend público el 21 de octubre. No son datos escritos a mano para llenar la interfaz: proceden del GET de Laravel, con device_id hydromate-01, contrato v3 y origen real en las variables mostradas.')
table(['Secuencia','Hora local','°C','pH','TDS ppm','Luz lux','Agua'],[
(r['sequence'],localtime(r['created_at']),f"{r['temperature_c']:.2f}" if r['temperature_c'] is not None else '—',f"{r['ph']:.2f}" if r['ph'] is not None else '—',f"{r['tds_ppm']:.2f}" if r['tds_ppm'] is not None else '—',f"{r['light_lux']:.2f}" if r['light_lux'] is not None else '—','Sí' if r['water_present'] else 'No') for r in records[:10]
],[.95,1.15,.65,.65,1.05,1.1,.45],10.5)
p('Fecha: 21 de octubre de 2026. Hora mostrada: America/Tijuana. En el backend las fechas permanecen en UTC. El archivo adjunto conserva 20 filas completas, incluidos id, secuencia, captured_at, created_at y sources.')
first=records[0]
p('Ejemplo de una adquisición guardada',bold=True)
code(json.dumps({k:first.get(k) for k in ('id','device_id','sequence','temperature_c','ph','tds_ppm','light_lux','water_present','captured_at')},ensure_ascii=False,indent=2))
p('Estos valores demuestran lectura, comunicación y persistencia. No se utilizan como prueba de calibración ni como recomendación de dosificación para el cultivo.')
p('Evidencia completa: evidencias/registros-publicos.json.',style='Caption')

page('La aplicación Android con datos reales')
p('Inicio presenta la lectura más reciente y su fecha. Historial reúne las recepciones guardadas y permite ver una gráfica sencilla. Ambas pantallas consultan al backend HTTPS; producción no sustituye fallos por datos de demostración.')
para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
para.add_run().add_picture(str(EV/'android-inicio.png'),width=Inches(2.75));para.add_run('   ')
para.add_run().add_picture(str(EV/'android-historial.png'),width=Inches(2.75))
p('Figuras 3 y 4. Capturas directas del teléfono el 21 de octubre. Inicio muestra TDS 772.77 ppm y 23.50 °C; Historial muestra registros reales con fecha y gráfica.',style='Caption')
p('La app maneja estados de carga, ausencia de lecturas y error de comunicación. Las variables opcionales pueden aparecer sin lectura; no se convierten en cero. Los estados están implementados en mobile/app/src/main/java/com/hydromate/mobile/ui/ReadingState.kt y comprobados por las pruebas Android.')
p('En esa misma base de código, Telemetry.kt consulta e interpreta la API; ui/dashboard/DashboardScreen.kt muestra la última lectura y ui/history/HistoryScreen.kt presenta los datos almacenados.',style='Caption')

page('Transmisiones consecutivas y tiempos')
p('En el ensayo físico público se registraron 14 muestras nuevas consecutivas, de la secuencia 4201 a la 4214. Todas recibieron STORED después de persistir. Hubo cero publicaciones fallidas y cero fallos MQTT después de conectar. Antes de sincronizar la cuenta se registró un rechazo de autenticación, que se conserva en la bitácora.')
table(['Secuencia','Hora de adquisición local','Luz lux','Confirmación'],[
(r['sequence'],localtime(r['captured_at']),r['light_lux'],'STORED') for r in physical['measurements']
],[1.05,2.45,1.25,2.25],10)
p('En ese primer ensayo TDS y temperatura todavía figuraban ausentes o inválidos; eso no elimina las mediciones físicas de luz, pH y flotador. Un ensayo posterior incorporó TDS y temperatura reales y quedó comparado con API y Android.')
table(['Tramo medido','Resultado','Cómo interpretarlo'],[
('SAMPLE a STORED','258–577 ms; media 409 ms','Incluye envío y confirmación de guardado; no incluye la pantalla.'),
('Adquisición a evidencia en Android','Hasta 27.14 s en una observación pública','Incluye espera de actualización y captura. Es una aproximación superior, no una latencia continua.'),
('Ensayo local anterior','8.24 s hasta verificar la interfaz','Antecedente local con apertura de app e inspección USB; no se presenta como medición de nube.')],[1.7,2.1,3.2],10)
p('Fuentes: docs/evidence/public-flow-esp32-android-2026-10-21.json y pruebas/latencia-android-reintento-2026-10-07.json.',style='Caption')

page('Pruebas de errores y correcciones')
p('La validación también contempla qué ocurre cuando algo falla. Los ensayos con mensajes incorrectos se hicieron en el entorno local aislado de pruebas; no se presentan como una inserción pública en la VPS.')
table(['Situación','Resultado observado','Corrección o manejo'],[
('pH incorrecto de 50','HTTP 422 y ACK rejected en el ensayo MQTT local.','Validación por campo; no se guarda la fila inválida.'),
('Reenvío idéntico','ACK stored sin otra fila para la misma identidad.','Índice único y comparación de contenido.'),
('Misma identidad con otro contenido','conflict; no sobrescribe la primera medición.','Detener el pendiente y reconciliar identidad.'),
('Red de 5 GHz o nombre incorrecto','ESP32 no encontraba la red.','Red de 2.4 GHz y configuración privada corregidas.'),
('Credenciales MQTT desalineadas','Broker rechazó autenticación.','Responsable sincronizó las cuentas; conexión y suscripción confirmadas.'),
('Corte y recuperación de conexión local','Secuencia 901 conservada y entregada después; 902 también guardada.','Reintento de la misma adquisición; evidencia de reconexión del 7 de octubre.'),
('Primer sensor TDS defectuoso','TDS no disponible en la etapa anterior.','Sensor reemplazado; entrada A1 habilitada y lectura real verificada.'),
('Temperatura inicialmente inválida','Campo null en la evidencia inicial.','Se recuperó lectura real; 23.50 °C observados junto con TDS.')],[1.55,2.7,2.75],10)
p('Respuesta comprensible del backend',bold=True)
p('Una lectura correcta devuelve HTTP 201 con el mensaje “Medicion registrada correctamente”. Una lectura fuera de rango devuelve HTTP 422 con errores del campo afectado. Un duplicado devuelve HTTP 409 con “Medicion duplicada”. El bridge transforma esos resultados en confirmaciones que el ESP32 puede distinguir.')
p('Se ejecutaron además 29 pruebas del backend con 194 comprobaciones, 10 pruebas Python y pruebas C++ de conversión y cambios. Android compiló en debug/release, pasó lint y ejecutó 31 pruebas JVM sin fallos; dos pruebas opcionales quedaron omitidas. Estas pruebas de software complementan la demostración física.')
p('Fuentes: docs/evidence/production-http-2026-10-14.json, production-mqtt-tls-2026-10-14.json, pruebas/esp32-reconexion-2026-10-07.json y docs/evidence/tds-events-2026-10-21.json.',style='Caption')

page('Cronograma hasta el 2 de diciembre')
p('La planeación parte del trabajo ya realizado y utiliza el 2 de diciembre de 2026 como fecha final indicada para el proyecto. Las etapas anteriores se conservan como antecedentes; las semanas restantes son una propuesta de trabajo, no resultados ya obtenidos.')
table(['Periodo transcurrido','Avance respaldado'],[
('20–22 septiembre','Compras de estructura e hidráulica registradas; revisión de cargos y alternativas de dosificación.'),
('Finales de septiembre','Prioridad definida para PostgreSQL, API Laravel, Android mínimo y ESP32 físico.'),
('3–7 octubre','Interfaz Android, sensores, contrato y ensayos locales; primeras diez entregas físicas.'),
('14–21 octubre','Empaquetado de producción, VPS rootless, HTTPS/MQTT TLS, integración pública y TDS de reemplazo.')],[1.6,5.4],10)
table(['Semana','Trabajo propuesto','Resultado esperado'],[
('21–25 oct','Cerrar evidencias del Taller 5 y revisar montaje.','Documento, fotos, video y requisitos trazados.'),
('26 oct–1 nov','Estabilizar sensores y revisar alimentación y cableado.','Bitácora de lecturas y conexiones reproducibles.'),
('2–8 nov','Calibrar pH/TDS y comparar mediciones repetidas.','Constantes y procedimiento de calibración documentados.'),
('9–15 nov','Preparar seguridad local y salida de banco para alfa.','Límites y estado seguro antes de cualquier acción física.'),
('16–22 nov','Integrar el actuador autorizado y probar con agua.','Comando y respuesta seguros, sin dosificación no validada.'),
('23–29 nov','Ensayo prolongado, fallos y documentación final.','Resultados de estabilidad y evidencia de recuperación.'),
('30 nov–2 dic','Revisión final, respaldo y demostración.','Paquete final y pendientes claramente identificados.')],[1.25,3.0,2.75],10)
p('El control remoto de actuadores se aplazó hasta el prototipo alfa, con autorización docente según la confirmación del responsable. La justificación técnica es validar primero sensores, límites, interbloqueos y estado seguro. No se afirma que los actuadores ya estén probados.')

page('Presupuesto actualizado')
p('Importes en MXN. El presupuesto se actualiza con las compras y el precio pagado confirmados por el responsable el 21 de octubre de 2026.')
p('Componentes del lote electrónico',bold=True)
table(['Componente','Cantidad','Precio total'],[('Kit de 840 cables jumper', '1 kit', '$154.00'), ('Sensor flotador', '1', '$65.00'), ('Módulo MOSFET dual Noyito', '1', '$60.00'), ('Sensor de pH PH014', '1', '$339.00'), ('Convertidor XL4015', '2', '$139.98'), ('Kit DS18B20 (2 sensores)', '1 kit', '$119.28'), ('BH1750 GY-302', '1', '$57.99'), ('Fuente de alimentación 24 V, 5 A', '1', '$162.75'), ('ADS1115', '1', '$76.23'), ('Suma del desglose', '', '$1,174.23')],[4.25,0.85,1.9],9.5)
p('Los precios desglosados suman $1,174.23, mientras que el importe pagado por el lote fue $1,168.00. La diferencia es de $6.23; no se informó el concepto del ajuste. Para el total del proyecto se utiliza el importe pagado, sin sumar nuevamente el desglose.')
p('Resumen del presupuesto',bold=True)
table(['Concepto','Costo'],[('Compras anteriores', '$2,055.94'), ('Electrónica completa (precio pagado)', '$1,168.00'), ('Segundo sensor TDS', '$270.00'), ('Total gastado confirmado', '$3,493.94'), ('Presupuesto máximo', '$6,000.00'), ('Disponible', '$2,506.06')],[5.0,2.0],10)
p('La VPS fue prestada y su costo para el proyecto es $0.00. El saldo disponible no representa una compra comprometida. Los importes de compras anteriores se conservan como un total confirmado, sin reconstruirlos con los subtotales históricos de septiembre.')
p('Fuente: relación de componentes, precios pagados y presupuesto máximo proporcionados por el responsable el 21 de octubre de 2026. No se presentan como comprobantes fiscales verificados.',style='Caption')

page('Riesgos y acciones previstas')
p('La matriz se concentra en lo que puede afectar la viabilidad y el siguiente prototipo. Las medidas futuras de control se revisan antes de conectar actuadores.')
table(['Riesgo','Impacto','Prevención','Contingencia'],[
('Wi-Fi o Internet no disponible','No llegan nuevas lecturas.','Verificar 2.4 GHz, señal y acceso antes del ensayo.','Conservar pendiente y reenviar al recuperar red.'),
('Sensor inválido o con deriva','La lectura deja de ser confiable.','Revisar cableado, filtrar y calibrar por separado.','Enviar null/invalid y reparar; no inventar valores.'),
('Reinicio o mensaje repetido','Duplicados o identidad en conflicto.','Secuencias NVS e índice único por torre.','Reconciliar identidad sin borrar la base.'),
('Interferencia entre servicios de VPS','Interrupción o acceso cruzado.','Rootless, cuentas, redes y volúmenes separados.','Intervenir solo HydroMate y recuperar su respaldo.'),
('Exposición de credenciales','Acceso no autorizado.','Secretos fuera de Git y revisión de capturas.','Rotar credencial afectada y revisar accesos.'),
('Lectura antigua entregada tarde','La app parece reciente sin serlo.','Conservar adquisición y recepción; calcular antigüedad.','Mostrar estado antiguo y revisar conexión.'),
('Actuador conectado sin protecciones','Daño al equipo o al cultivo.','Aplazar hasta alfa; límites y estado seguro primero.','Desenergizar y probar con agua en banco.'),
('Gasto o material sin confirmar','Retraso o compra duplicada.','Inventario, recibos y presupuesto actualizado.','Priorizar el flujo mínimo y cotizar el faltante.')],[1.25,1.3,2.25,2.2],10)
p('La prioridad para el siguiente incremento es estabilizar el montaje y documentar calibración. La prueba actual no habilita automáticamente control químico ni dosificación.')

page('Revisión de requisitos técnicos')
p('C significa comprobado con ejecución o evidencia física. V significa verificado en código o en un ensayo local de software; el alcance se identifica en la evidencia. La aprobación de viabilidad corresponde al docente.')
table(['Requisito del taller','Estado','Dónde se demuestra'],[
('Dato real de extremo a extremo','C','Fig. 1; 14 muestras públicas y capturas Android.'),
('Leer sensor o entrada física','C','DS18B20, BH1750, flotador, pH y TDS.'),
('Procesamiento básico','C','Firmware: mediana, conversiones y validación.'),
('Conectarse a Internet','C','ESP32 conectado por Wi-Fi y MQTT TLS público.'),
('Enviar un tipo de dato al backend','C','Mensajes v3 y confirmación STORED.'),
('Identificar medición','C','device_id, sequence, captured_at y created_at.'),
('Backend desplegado en nube','C','API HTTPS pública y /up HTTP 200.'),
('Recibir dato del embebido','C','Serial y API coinciden por secuencia.'),
('Validar campos indispensables','V','StoreMeasurementRequest y ensayos 201/422/409.'),
('Registrar en base de datos','C','STORED y recuperación posterior por GET.'),
('Consulta para frontend','C','GET historial/última/estado e interfaz Android.'),
('Respuesta válida e incorrecta','V','Prueba HTTP y MQTT local con ph=50, HTTP 422.'),
('Una tabla o colección','C','Tabla measurements y migraciones.'),
('Identificador de registro','C','id y clave única device_id/sequence.'),
('Variable y valor','C','Columnas de sensores y ejemplo JSON.'),
('Unidad cuando corresponda','C','°C, ppm, lux y µS/cm en contrato; pH/nivel sin unidad.'),
('Fecha y hora','C','UTC en API y hora local en pantalla.'),
('Mínimo diez registros reales','C','Tabla de diez filas; archivo con 20 registros públicos.'),
('Consulta de recuperación','C','SQL documentado y GET funcional.'),
('Pantalla funcional','C','Inicio e Historial en teléfono físico.'),
('Comunicación real con backend','C','Android producción consume HTTPS público.'),
('Dato reciente y fecha','C','Captura de Inicio.'),
('Lista o gráfica de almacenados','C','Captura de Historial.'),
('Carga vacío o error','V','ReadingState y pruebas Android; sin fallback de producción.')],[2.75,.55,3.7],9.5)

page('Revisión de pruebas y entregables')
photo_state='C' if photos else 'P'
table(['Requisito del taller','Estado','Evidencia o situación'],[
('Diez transmisiones consecutivas','C','14 muestras físicas públicas 4201–4214.'),
('Registro de éxitos y fallos','C','14 STORED, cero publish fallidos; rechazo previo documentado.'),
('Tiempo aproximado sensor a frontend','C','Observación pública hasta 27.14 s, con alcance explicado.'),
('Dato incorrecto o pérdida temporal','V','ph=50 rechazado y reconexión física local 901–902.'),
('Problemas y correcciones','C','Bitácora resumida y archivos de pruebas.'),
('Código fuente inicial','C','embedded/, scripts/, hydromate-backend/ y mobile/.'),
('Diagrama de conexión y flujo','C','Figuras 1–2; fuentes SVG y PNG adjuntas.'),
('Video breve de funcionamiento','C','evidencias/android-datos-reales.mp4, pantalla conectada a producción.'),
('Evidencia de backend desplegado','C','Consultas públicas 200 con fecha y registros reales.'),
('Estructura inicial de BD','C','Migraciones, tabla y exportación de esquema local.'),
('Frontend con datos reales','C','Capturas directas de Inicio e Historial.'),
('Git y README básico','C','github.com/komvo/HydroMate; README raíz y guías por componente.'),
('Cronograma restante','C','Plan de semanas hasta el 2 de diciembre.'),
('Presupuesto preliminar','C','Gasto $3,493.94; máximo $6,000.00; disponible $2,506.06. Diferencia de desglose documentada.'),
('Al menos cinco riesgos con acciones','C','Ocho riesgos con prevención y contingencia.'),
('Actuador si es función central','Aut.','Aplazamiento hasta alfa informado como autorizado por docente.'),
('Fotografías de sensores y ESP32',photo_state,'Anexo de montaje.' if photos else 'Pendientes de incorporar; complementan la evidencia del armado.')],[2.75,.55,3.7],9.5)
p(('Aut. indica aplazamiento autorizado según la confirmación del responsable. ' if photos else 'P indica anexo pendiente y Aut. indica aplazamiento autorizado según la confirmación del responsable. ') + 'El presupuesto se actualiza con importes pagados confirmados; la diferencia del desglose electrónico se señala en su sección.')
p('Para este taller no se requiere autenticación completa, lógica de negocio definitiva, todas las pantallas, diseño adaptable terminado, documentación completa de API, calibración ni pruebas unitarias formales. El desarrollo siguiente conserva el proyecto aprobado; cualquier sustitución requiere autorización expresa del docente.')

page('Guía de evidencias y demostración')
table(['Evidencia','Qué permite revisar'],[
('evidencias/android-inicio.png','Última recepción y cinco variables físicas.'),
('evidencias/android-historial.png','Gráfica y recepciones almacenadas.'),
('evidencias/android-datos-reales.mp4','Funcionamiento de la app conectada al backend público.'),
('evidencias/api-publica.json','Comprobaciones HTTPS de salud, historial y estado.'),
('evidencias/registros-publicos.json','Veinte registros obtenidos del backend público.'),
('diagramas/flujo.svg y conexiones.svg','Fuentes editables de los diagramas implementados.'),
('docs/evidence/public-flow-esp32-android-2026-10-21.json','Catorce transmisiones físicas y tiempos de confirmación.'),
('docs/evidence/tds-events-2026-10-21.json','TDS real, cambios de luz y nivel, y comparación con API.'),
('talleres/taller-05/pruebas/','Ensayos anteriores de esquema, reconexión y errores.'),
('docs/API.md y TELEMETRY_CONTRACT.md','Campos, unidades, estados y formato del mensaje.')],[3.6,3.4],9.5)
p('Cómo mostrarlo durante la revisión',bold=True)
p('1. Enseñar el ESP32 y los sensores conectados. Mantener la red de 2.4 GHz y abrir HydroMate con el perfil Producción y dispositivo hydromate-01.')
p('2. Cubrir y destapar el sensor de luz o mover el flotador. Mostrar que la fecha y los valores cambian en Inicio y que la nueva recepción aparece en Historial.')
p('3. Abrir el GET del backend y comparar la secuencia, fecha y valores con la evidencia. Mostrar las diez filas y el registro de STORED; explicar el rechazo de datos incorrectos con la prueba guardada.')
p('4. Cerrar con el cronograma, presupuesto y riesgos. La actuación física se revisará en alfa conforme al aplazamiento informado.')
p('Repositorio: https://github.com/komvo/HydroMate\nVersión del incremento de sensores: bfda030.\nLas rutas de código y evidencia indicadas sin prefijo parten de la raíz del repositorio.')

if photos:
    for item in photos:
        page(item['title'])
        p(item['intro'])
        figure(HERE/item['path'],item['caption'],float(item.get('width',5.9)))

output=HERE/'HydroMate_Taller_5.docx'
doc.save(output)
manifest={'date':'2026-10-21','document':output.name,'photos_incorporated':len(photos),'public_real_records':len(records),'physical_consecutive_confirmed':physical['new_samples_confirmed_stored'],'video':str((EV/'android-datos-reales.mp4').relative_to(HERE)),'scope':'Taller 5 minimum viability evidence','authoring':'python-docx; Letter portrait; visual review required'}
(HERE/'indice-evidencias.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False)+'\n',encoding='utf8')
print(output)
